"""
DNA Processor — extrae texto de archivos subidos y llama a Claude para analizar el canal.
Soporta: PDF, DOCX, XLSX, TXT
"""
import io
import json
import os
import re
import tempfile
from typing import Optional

import anthropic

# ── Extracción de texto por formato ────────────────────────────────────────

def _extract_txt(data: bytes) -> str:
    for enc in ("utf-8", "latin-1", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _extract_pdf(data: bytes) -> str:
    try:
        import pdfplumber
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            pages = [p.extract_text() or "" for p in pdf.pages]
        return "\n".join(pages)
    except Exception as e:
        print(f"[dna] pdfplumber failed: {e}")
        return ""


def _extract_docx(data: bytes) -> str:
    try:
        from docx import Document
        doc = Document(io.BytesIO(data))
        return "\n".join(p.text for p in doc.paragraphs)
    except Exception as e:
        print(f"[dna] docx failed: {e}")
        return ""


def _extract_xlsx(data: bytes) -> str:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
        rows = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) if c is not None else "" for c in row]
                rows.append("\t".join(cells))
        return "\n".join(rows)
    except Exception as e:
        print(f"[dna] xlsx failed: {e}")
        return ""


def _extract_csv(data: bytes) -> str:
    return _extract_txt(data)


def extract_text_from_file(filename: str, data: bytes) -> str:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext == "pdf":
        return _extract_pdf(data)
    if ext in ("doc", "docx"):
        return _extract_docx(data)
    if ext in ("xlsx", "xls"):
        return _extract_xlsx(data)
    if ext == "csv":
        return _extract_csv(data)
    return _extract_txt(data)


# ── Prompt de análisis del DNA ──────────────────────────────────────────────

_DNA_SYSTEM = """Eres un analista experto en canales de YouTube especializado en contenido de meditación, manifestación y desarrollo personal."""

_DNA_USER_TMPL = """Analiza todos los documentos proporcionados y extrae el DNA del canal.

DOCUMENTOS:
{docs}

Extrae:
1. VOZ Y ESTILO: muletillas frecuentes, frases de apertura típicas, frases de cierre, tono dominante, vocabulario propio del nicho, palabras que aparecen en los mejores guiones.
2. PATRONES DE TÍTULOS EXITOSOS: estructura, longitud, palabras de alto CTR, fórmulas que se repiten en los títulos con más vistas. Genera 5-8 plantillas de título basadas en los patrones reales.
3. ESTRUCTURA DE GUIÓN GANADORA: cómo abren los guiones exitosos, cómo está estructurado el cuerpo, cómo cierran. Incluye tiempos aproximados si están disponibles.
4. LISTA NEGRA: frases, palabras, estructuras y tonos que aparecen en los guiones de bajo rendimiento. Qué NO debe hacer el canal.
5. AFIRMACIONES MODELO: si el canal usa afirmaciones, extrae las que mejor funcionaron y define sus características (longitud, tiempo verbal, uso de metáforas, etc.)
6. FÓRMULA DE COMBINACIÓN ÓPTIMA: basado en los datos, cuál es la combinación de tono + estructura + dirección energética que maximiza el rendimiento.

Responde SOLO en JSON válido con esta estructura exacta:
{{
  "voz_estilo": {{
    "muletillas": [],
    "apertura_tipica": "",
    "cierre_tipico": "",
    "tono_dominante": "",
    "vocabulario_nicho": [],
    "palabras_clave": []
  }},
  "patrones_titulos": {{
    "plantillas": [],
    "palabras_alto_ctr": [],
    "longitud_optima": "",
    "formulas": []
  }},
  "estructura_guion": {{
    "hook_segundos": "",
    "estructura_cuerpo": [],
    "cierre": "",
    "duracion_total": ""
  }},
  "lista_negra": {{
    "frases_prohibidas": [],
    "tonos_evitar": [],
    "estructuras_evitar": []
  }},
  "afirmaciones_modelo": {{
    "ejemplos": [],
    "reglas": []
  }},
  "formula_optima": {{
    "tono": "",
    "estructura": "",
    "direccion": "",
    "marco": "",
    "rendimiento_esperado": ""
  }},
  "resumen_canal": ""
}}"""


def _build_docs_block(files: list[dict]) -> str:
    """files: list of {zone, filename, text}"""
    ZONE_LABELS = {
        "exitosos":     "GUIONES EXITOSOS",
        "titulos":      "HISTORIAL DE TÍTULOS",
        "analisis":     "ANÁLISIS DEL CANAL",
        "bajo_rendimiento": "GUIONES DE BAJO RENDIMIENTO",
    }
    blocks = []
    for f in files:
        label = ZONE_LABELS.get(f["zone"], f["zone"].upper())
        blocks.append(f"=== {label}: {f['filename']} ===\n{f['text'][:8000]}")
    return "\n\n".join(blocks)


# ── Claude API call ─────────────────────────────────────────────────────────

def process_dna_with_claude(files: list[dict], api_key: str) -> dict:
    """
    files: list of {zone, filename, text}
    Returns the parsed DNA dict.
    Raises on failure.
    """
    client = anthropic.Anthropic(api_key=api_key)

    docs_block = _build_docs_block(files)
    user_prompt = _DNA_USER_TMPL.format(docs=docs_block)

    response = client.messages.create(
        model="claude-sonnet-4-6",
        max_tokens=16000,
        system=_DNA_SYSTEM,
        messages=[{"role": "user", "content": user_prompt}],
    )

    # Detectar respuesta truncada antes de parsear
    if response.stop_reason == "max_tokens":
        raise ValueError(
            "La respuesta de Claude fue cortada porque los archivos son muy grandes. "
            "Reduce el número de archivos o usa documentos más cortos."
        )

    raw = response.content[0].text.strip()

    # Strip markdown fences if present
    raw = re.sub(r"^```(?:json)?\s*", "", raw)
    raw = re.sub(r"\s*```$", "", raw)

    try:
        return json.loads(raw)
    except json.JSONDecodeError as e:
        raise ValueError(f"Claude devolvió JSON inválido: {e}. Respuesta parcial: {raw[:200]}")


# ── Persistencia en disco ───────────────────────────────────────────────────

DNA_DIR = os.getenv("DNA_STORAGE_PATH", os.path.join(os.path.dirname(__file__), "..", "data", "dna"))


def _dna_path(user_id: int) -> str:
    os.makedirs(DNA_DIR, exist_ok=True)
    return os.path.join(DNA_DIR, f"user_{user_id}.json")


def save_dna(user_id: int, dna: dict) -> None:
    with open(_dna_path(user_id), "w", encoding="utf-8") as f:
        json.dump(dna, f, ensure_ascii=False, indent=2)


def load_dna(user_id: int) -> Optional[dict]:
    path = _dna_path(user_id)
    if not os.path.exists(path):
        return None
    try:
        with open(path, encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return None


def delete_dna(user_id: int) -> bool:
    path = _dna_path(user_id)
    if os.path.exists(path):
        os.remove(path)
        return True
    return False
